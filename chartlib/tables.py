"""Append three-line tables while inheriting typography from the Word document."""
from __future__ import annotations

import os
import re

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt

FONT_CN = None
FONT_EN = None
SZ_TITLE = None
SZ_BODY = None
TABLE_PREFIX = "\u8868"
CHINESE_FULL_STOP = "\u3002"


def _set_run(run, size_pt: float | None, bold: bool = False, italic: bool = False):
    if FONT_EN:
        run.font.name = FONT_EN
    if size_pt is not None:
        run.font.size = Pt(size_pt)
    run.font.bold = bold
    run.font.italic = italic
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    if FONT_EN:
        rfonts.set(qn("w:ascii"), FONT_EN)
        rfonts.set(qn("w:hAnsi"), FONT_EN)
    if FONT_CN:
        rfonts.set(qn("w:eastAsia"), FONT_CN)


def _add_para(doc, text: str, size_pt: float | None, align=WD_ALIGN_PARAGRAPH.LEFT,
              indent: bool = False, line_spacing: float = 1.25) -> None:
    p = doc.add_paragraph()
    p.alignment = align
    pf = p.paragraph_format
    pf.line_spacing = line_spacing
    pf.space_before = Pt(0)
    pf.space_after = Pt(0)
    if indent:
        if size_pt is not None:
            pf.first_line_indent = Pt(size_pt * 2)
    _set_run(p.add_run(text), size_pt)


def _set_cell(cell, text, size_pt: float | None, bold: bool = False, italic: bool = False,
              center: bool = True):
    cell.text = ""
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER if center else WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.line_spacing = 1.0
    _set_run(p.add_run(str(text)), size_pt, bold=bold, italic=italic)


def _border(elm_tag: str, sz_eighth_pt: int):
    el = OxmlElement(elm_tag)
    el.set(qn("w:val"), "single")
    el.set(qn("w:sz"), str(sz_eighth_pt))
    el.set(qn("w:space"), "0")
    el.set(qn("w:color"), "000000")
    return el


def _apply_three_line(table) -> None:
    """Apply 1.5/0.75/1.5 pt horizontal rules and remove vertical rules."""
    tbl = table._tbl
    tblPr = tbl.tblPr
    # Replace all existing table borders.
    for old in tblPr.findall(qn("w:tblBorders")):
        tblPr.remove(old)
    borders = OxmlElement("w:tblBorders")
    borders.append(_border("w:top", 12))
    borders.append(_border("w:bottom", 12))
    for side in ("w:left", "w:right", "w:insideH", "w:insideV"):
        el = OxmlElement(side)
        el.set(qn("w:val"), "none")
        borders.append(el)
    tblPr.append(borders)
    # Add the thin rule below every header cell.
    for cell in table.rows[0].cells:
        tcPr = cell._tc.get_or_add_tcPr()
        for old in tcPr.findall(qn("w:tcBorders")):
            tcPr.remove(old)
        tcb = OxmlElement("w:tcBorders")
        tcb.append(_border("w:bottom", 6))
        tcPr.append(tcb)


def _existing_numbers(doc_path: str) -> dict[int, list[int]]:
    """Return existing table sequence numbers by chapter."""
    used: dict[int, list[int]] = {}
    if not os.path.exists(doc_path):
        return used
    try:
        doc = Document(doc_path)
    except Exception:
        return used
    pat = re.compile(TABLE_PREFIX + r"\s*(\d+)\s*[.]\s*(\d+)")
    for p in doc.paragraphs:
        m = pat.match(p.text.strip())
        if m:
            used.setdefault(int(m.group(1)), []).append(int(m.group(2)))
    return used


def next_number(doc_path: str, chapter: int) -> int:
    used = _existing_numbers(doc_path)
    seq = [n for n in used.get(chapter, []) if n > 0]
    return (max(seq) + 1) if seq else 1


class TableSpec:
    """Specification for one three-line table."""

    def __init__(self, chapter: int, title: str, header: list[str], rows: list[list],
                 note: list[str] | None = None, source: str | None = None,
                 italic_cols: list[int] | None = None):
        if not title or CHINESE_FULL_STOP in title or "." in title:
            raise ValueError("The table title must not contain terminal punctuation.")
        if not header or not rows:
            raise ValueError("The table header and body cannot be empty.")
        if not all(len(r) == len(header) for r in rows):
            raise ValueError("Every body row must have the same width as the header.")
        self.chapter = int(chapter)
        self.title = title
        self.header = header
        self.rows = rows
        self.note = note or []
        self.source = source
        self.italic_cols = set(italic_cols or [])


def _append_table(doc: Document, spec: TableSpec, number: str) -> None:
    # Center the title and keep the number and title on one line.
    _add_para(doc, f"{TABLE_PREFIX}{spec.chapter}.{number} {spec.title}",
              SZ_TITLE, align=WD_ALIGN_PARAGRAPH.CENTER)
    table = doc.add_table(rows=1 + len(spec.rows), cols=len(spec.header))
    table.autofit = True
    # Header row.
    for j, h in enumerate(spec.header):
        _set_cell(table.rows[0].cells[j], h, SZ_BODY, bold=False)
    # Body rows.
    for i, row in enumerate(spec.rows, start=1):
        for j, val in enumerate(row):
            italic = j in spec.italic_cols
            _set_cell(table.rows[i].cells[j], val, SZ_BODY, italic=italic)
    _apply_three_line(table)
    # Add a small spacer before notes or the next table.
    _add_para(doc, "", 6)
    # Add user-approved notes in the supplied order.
    for line in spec.note:
        _add_para(doc, line, SZ_TITLE, align=WD_ALIGN_PARAGRAPH.LEFT)
    if spec.source:
        _add_para(doc, spec.source, SZ_TITLE, align=WD_ALIGN_PARAGRAPH.LEFT)


def render_tables(specs: list[TableSpec], doc_path: str) -> str:
    """Append tables and continue chapter numbering across calls."""
    doc = None
    if os.path.exists(doc_path):
        try:
            doc = Document(doc_path)
        except Exception:
            doc = None
    if doc is None:
        doc = Document()
    used = _existing_numbers(doc_path)
    made: list[str] = []
    for spec in specs:
        seq = used.setdefault(spec.chapter, [])
        n = (max(seq) + 1) if seq else 1
        seq.append(n)
        _append_table(doc, spec, n)
        made.append(f"{TABLE_PREFIX}{spec.chapter}.{n} {spec.title}")
    doc.save(doc_path)
    return "\n".join(made)
